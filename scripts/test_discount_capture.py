import openpyxl
import re
import sys
import unicodedata

sys.stdout.reconfigure(encoding='utf-8')

def norm(s):
    if not s: return ''
    s = unicodedata.normalize('NFKD', str(s)).encode('ASCII', 'ignore').decode('utf-8').lower()
    return re.sub(r'[^a-z0-9]', '', s)

wb = openpyxl.load_workbook('Planilha de Dados - Colégio Rodin 2026.xlsx', data_only=True)

sheets = ['6º ano - 2026', '7º ano - 2026', '8º ano - 2026', '9º ano - 2026 ', '1ª série - 2026', '2ª Série - 2026', '3ª Série - 2026']

total_students = 0
total_lp = 0
total_with_desc = 0
total_bolsa_100 = 0
discounts_summary = {}

for sname in sheets:
    ws = wb[sname]
    header_idx = {}
    for c in range(1, 130):
        v = ws.cell(1, c).value
        if v:
            header_idx[norm(v)] = c
    
    desc_col = header_idx.get(norm('Descrição do Desconto')) or 29
    origem_col = header_idx.get(norm('Escola de Origem'))
    
    for r in range(2, ws.max_row + 1):
        coc = ws.cell(r, 1).value
        name = ws.cell(r, 2).value
        if not coc or not name or len(str(name).strip()) < 3:
            continue
        
        total_students += 1
        d_text = str(ws.cell(r, desc_col).value or '').strip()
        origem = str(ws.cell(r, origem_col).value or '').strip() if origem_col else ''
        
        is_lp = bool('le perini' in d_text.lower() or 'dlp' in d_text.lower() or 'leperini' in d_text.lower() or 'le perini' in origem.lower())
        if is_lp:
            total_lp += 1
            
        is_100 = bool('100%' in d_text or 'bolsa 100' in d_text.lower() or 'integral' in d_text.lower())
        if is_100:
            total_bolsa_100 += 1
            
        m = re.search(r'(\d+)\s*%', d_text)
        pct = int(m.group(1)) if m else (100 if is_100 else (25 if is_lp else 0))
        if pct > 0:
            total_with_desc += 1
            discounts_summary[pct] = discounts_summary.get(pct, 0) + 1

print(f"Total Alunos: {total_students}")
print(f"Total Le Perini: {total_lp}")
print(f"Total com Desconto (>0%): {total_with_desc}")
print(f"Total Bolsa 100%: {total_bolsa_100}")
print("Distribuição dos Descontos (%):", sorted(discounts_summary.items()))
