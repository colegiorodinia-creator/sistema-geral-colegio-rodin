import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('Planilha de Dados - Colégio Rodin 2026.xlsx', data_only=True)
for sname in wb.sheetnames[:7]:
    ws = wb[sname]
    cols = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        if v:
            cols.append((c, str(v).strip()))
    print(f"Sheet: [{sname}] - Total colunas com header: {len(cols)}")
    discount_cols = [x for x in cols if any(k in x[1].lower() for k in ['desc', 'bolsa', 'origem', 'val', 'anu', 'obs'])]
    print("  Colunas financeiras/desconto encontradas:", discount_cols)
