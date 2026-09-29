import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('Planilha de Dados - Colégio Rodin 2026.xlsx', data_only=True)
for sname in wb.sheetnames[:7]:
    ws = wb[sname]
    found = 0
    print(f"\n--- Aba: [{sname}] ---")
    for r in range(2, min(ws.max_row + 1, 60)):
        name = ws.cell(r, 2).value
        v_desc = ws.cell(r, 27).value
        d_text = ws.cell(r, 29).value
        v_tot = ws.cell(r, 30).value
        if d_text or v_desc or v_tot:
            print(f"Row {r} | Nome: {name} | DescText: '{d_text}' | VlDesc: {v_desc} | Tot: {v_tot}")
            found += 1
            if found >= 5:
                break
    if found == 0:
        print("  Nenhum valor encontrado nas primeiras 60 linhas.")
