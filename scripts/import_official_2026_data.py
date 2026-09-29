import openpyxl
import json
import re
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

# Tabelas Oficiais 2027 do Colégio Rodin
RATES_2027 = {
    '7º Ano EF': {
        'cycle': 'Ensino Fundamental',
        'nominal': 34663.20,
        'material_total': 5248.80,
        'material_extenso': 'Cinco mil, duzentos e quarenta e oito reais e oitenta centavos',
        'material_parc': 437.40,
        'material_parc_extenso': 'Quatrocentos e trinta e sete reais e quarenta centavos',
        'material_parcelas': 12
    },
    '8º Ano EF': {
        'cycle': 'Ensino Fundamental',
        'nominal': 34663.20,
        'material_total': 5248.80,
        'material_extenso': 'Cinco mil, duzentos e quarenta e oito reais e oitenta centavos',
        'material_parc': 437.40,
        'material_parc_extenso': 'Quatrocentos e trinta e sete reais e quarenta centavos',
        'material_parcelas': 12
    },
    '9º Ano EF': {
        'cycle': 'Ensino Fundamental',
        'nominal': 34663.20,
        'material_total': 5248.80,
        'material_extenso': 'Cinco mil, duzentos e quarenta e oito reais e oitenta centavos',
        'material_parc': 437.40,
        'material_parc_extenso': 'Quatrocentos e trinta e sete reais e quarenta centavos',
        'material_parcelas': 12
    },
    '1ª Série EM': {
        'cycle': 'Ensino Médio',
        'nominal': 37752.00,
        'material_total': 5338.20,
        'material_extenso': 'Cinco mil, trezentos e trinta e oito reais e vinte centavos',
        'material_parc': 444.85,
        'material_parc_extenso': 'Quatrocentos e quarenta e quatro reais e oitenta e cinco centavos',
        'material_parcelas': 12
    },
    '2ª Série EM': {
        'cycle': 'Ensino Médio',
        'nominal': 37752.00,
        'material_total': 5338.20,
        'material_extenso': 'Cinco mil, trezentos e trinta e oito reais e vinte centavos',
        'material_parc': 444.85,
        'material_parc_extenso': 'Quatrocentos e quarenta e quatro reais e oitenta e cinco centavos',
        'material_parcelas': 12
    },
    '3ª Série EM': {
        'cycle': 'Ensino Médio',
        'nominal': 43243.20,
        'material_total': 7575.60,
        'material_extenso': 'Sete mil, quinhentos e setenta e cinco reais e sessenta centavos',
        'material_parc': 631.30,
        'material_parc_extenso': 'Seiscentos e trinta e um reais e trinta centavos',
        'material_parcelas': 12
    },
    'Concluinte': {
        'cycle': 'Ensino Médio',
        'nominal': 0.0,
        'material_total': 0.0,
        'material_extenso': '',
        'material_parc': 0.0,
        'material_parc_extenso': '',
        'material_parcelas': 0
    }
}

GRADE_PROGRESSION = {
    '6º ano - 2026': ('6º Ano EF', '7º Ano EF', 'Ensino Fundamental'),
    '7º ano - 2026': ('7º Ano EF', '8º Ano EF', 'Ensino Fundamental'),
    '8º ano - 2026': ('8º Ano EF', '9º Ano EF', 'Ensino Fundamental'),
    '9º ano - 2026 ': ('9º Ano EF', '1ª Série EM', 'Ensino Fundamental'),
    '1ª série - 2026': ('1ª Série EM', '2ª Série EM', 'Ensino Médio'),
    '2ª Série - 2026': ('2ª Série EM', '3ª Série EM', 'Ensino Médio'),
    '3ª Série - 2026': ('3ª Série EM', 'Concluinte', 'Ensino Médio')
}

def clean_val(v):
    if v is None:
        return ''
    s = str(v).strip()
    return '' if s in ['None', 'null', '—', '-'] else s

def clean_money(v):
    if not v:
        return 0.0
    try:
        s = str(v).replace('R$', '').replace(' ', '').replace('.', '').replace(',', '.')
        return float(s)
    except:
        return 0.0

def clean_date(v):
    if not v:
        return ''
    s = str(v).strip()
    # Check if standard date string
    m = re.match(r'(\d{2})/(\d{2})/(\d{4})', s)
    if m:
        return f"{m.group(1)}/{m.group(2)}/{m.group(3)}"
    m2 = re.match(r'(\d{4})-(\d{2})-(\d{2})', s)
    if m2:
        return f"{m2.group(3)}/{m2.group(2)}/{m2.group(1)}"
    return s

def split_address(full_logr):
    if not full_logr:
        return ('', '', '')
    parts = full_logr.split(',')
    street = parts[0].strip()
    number = ''
    comp = ''
    if len(parts) > 1:
        rest = ','.join(parts[1:]).strip()
        m = re.match(r'^(\d+[A-Za-z]?)\s*(.*)$', rest)
        if m:
            number = m.group(1).strip()
            comp = m.group(2).strip()
        else:
            number = rest
    return (street, number, comp)

def calculate_installments(gross, nominal, is_le_perini):
    if is_le_perini:
        regular = round(gross / 13.0, 2)
        return regular, regular
    else:
        first = round(nominal / 13.0, 2)
        remaining = max(0.0, gross - first)
        reg = round(remaining / 12.0, 2)
        return first, reg

def parse_student_row(ws, r, header_idx, current_grade, target_grade, cycle):
    def get_h(hname):
        idx = header_idx.get(hname.lower())
        if idx:
            return clean_val(ws.cell(row=r, column=idx).value)
        return ''

    rm = get_h('num_coc')
    name = get_h('nome_aluno')
    if not name:
        return None

    # Sexo estudante
    sex_m = get_h('sex_mas_est').lower()
    sex_f = get_h('sexo_fem_est').lower()
    student_gender = 'Fem.' if 'x' in sex_f else 'Masc.'

    # Sexo responsável
    resp_sex_m = get_h('sexo_mas_resp').lower()
    resp_sex_f = get_h('sexo_fem_resp').lower()
    resp_gender = 'Fem.' if 'x' in resp_sex_f else 'Masc.'

    # Desconto 2026 e Descrição Oficial da Planilha
    desc_desc = get_h('Descrição do Desconto') or get_h('descricao do desconto') or get_h('desc_desc')
    if not desc_desc and ws.cell(row=r, column=29).value:
        desc_desc = clean_val(ws.cell(row=r, column=29).value)

    vl_anuid_desc = clean_money(get_h('vl_anuidade_desc') or ws.cell(row=r, column=27).value)
    anuid_total = clean_money(get_h('anuidade_total') or ws.cell(row=r, column=30).value)
    origem = get_h('Escola de Origem') or get_h('escola de origem')

    is_le_perini = bool(
        'le perini' in desc_desc.lower() or 
        'dlp' in desc_desc.lower() or 
        'leperini' in desc_desc.lower() or
        'le perini' in origem.lower() or
        'dlp' in origem.lower()
    )

    # Identificar se é Bolsa 100% / Integral
    is_100 = bool(
        '100%' in desc_desc or 
        'bolsa 100' in desc_desc.lower() or 
        'integral' in desc_desc.lower() or
        'bolsa func' in desc_desc.lower() or
        'permuta' in desc_desc.lower()
    )

    # Identificar percentual de desconto
    discount_pct = 0.0
    m_pct = re.search(r'(\d+)\s*%', desc_desc)
    if is_100:
        discount_pct = 1.0
    elif m_pct:
        discount_pct = float(m_pct.group(1)) / 100.0
    elif anuid_total > 0 and vl_anuid_desc > 0 and vl_anuid_desc < anuid_total:
        discount_pct = round(1.0 - (vl_anuid_desc / anuid_total), 2)
    elif is_le_perini:
        discount_pct = 0.25

    # Endereço
    logr_resp = get_h('Logr_resp') or get_h('logradouro')
    street, num, comp = split_address(logr_resp)

    # Turma atual 2026
    turma_2026 = get_h('turma') or 'A'
    turno = get_h('turno') or 'Manhã'

    # Se concluinte (3ª série 2026)
    is_concluinte = (target_grade == 'Concluinte')
    rates = RATES_2027[target_grade]

    if is_concluinte:
        nominal_2027 = 0.0
        gross_2027 = 0.0
        first_inst = 0.0
        reg_inst = 0.0
        status_rematricula = 'Concluiu'
        status_contrato = 'Concluiu'
        status_material = 'Concluiu'
    else:
        nominal_2027 = rates['nominal']
        if is_100 or discount_pct >= 1.0:
            gross_2027 = 0.0
            discount_pct = 1.0
        elif discount_pct > 0:
            gross_2027 = round(nominal_2027 * (1.0 - discount_pct), 2)
        else:
            gross_2027 = nominal_2027
        
        first_inst, reg_inst = calculate_installments(gross_2027, nominal_2027, is_le_perini)
        status_rematricula = 'Pendente'
        status_contrato = 'Pendente'
        status_material = 'Pendente'

    obs_desc = desc_desc if desc_desc else ('Sem desconto' if discount_pct == 0 else f'Desconto de {int(discount_pct*100)}% na anuidade')

    # Celulares com DDD
    ddd_est = get_h('ddd_cel_est')
    cel_est = get_h('cel_est')
    phone_est = f"({ddd_est}) {cel_est}" if ddd_est and cel_est else cel_est

    ddd_resp = get_h('ddd_cel_rep_fin')
    cel_resp = get_h('cel_rep_fin')
    phone_resp = f"({ddd_resp}) {cel_resp}" if ddd_resp and cel_resp else cel_resp

    ddd_pai = get_h('ddd_cel_pai')
    cel_pai = get_h('tel_cel_pai')
    phone_pai = f"({ddd_pai}) {cel_pai}" if ddd_pai and cel_pai else cel_pai

    ddd_mae = get_h('ddd_cel_mae')
    cel_mae = get_h('tel_cel_mae')
    phone_mae = f"({ddd_mae}) {cel_mae}" if ddd_mae and cel_mae else cel_mae

    record = {
        'rm_numero': rm,
        'coc_code': rm,
        'ano_letivo': 2027,
        'nome_completo_estudante': name,
        'serie_ano_atual': current_grade,
        'nova_serie_ano_2027': target_grade,
        'nivel_ensino': cycle,
        'periodo_turno': turno,
        'turma': turma_2026,
        'sexo_estudante': student_gender,
        'data_nascimento_estudante': clean_date(get_h('data_nasc_est')),
        'naturalidade_cidade': get_h('loc_nasc_est') or 'Indaiatuba',
        'naturalidade_uf': get_h('ocup_est') or 'SP',
        'nacionalidade_estudante': get_h('nacion_est') or 'Brasileira',
        'rg_estudante': get_h('rg_est'),
        'orgao_emissor_rg_estudante': get_h('org_exp_est') or 'SSP/SP',
        'data_emissao_rg_estudante': clean_date(get_h('data_emissao_est')),
        'cpf_estudante': get_h('cpf_est'),
        'celular_whatsapp_estudante': phone_est,
        'nome_responsavel_financeiro': get_h('Nome_rep_fin') or get_h('nome_mae') or get_h('nome_pai') or name,
        'parentesco_responsavel': get_h('parentesco') or 'Responsável',
        'sexo_responsavel': resp_gender,
        'data_nascimento_responsavel': clean_date(get_h('data_nasc_resp')),
        'profissao_responsavel': get_h('ocup_resp'),
        'estado_civil_responsavel': get_h('est_civil_res') or 'Casado(a)',
        'nacionalidade_responsavel': get_h('nacio_resp') or 'Brasileira',
        'rg_responsavel': get_h('rg_resp'),
        'orgao_emissor_rg_responsavel': get_h('org_exp_resp') or 'SSP/SP',
        'cpf_responsavel_financeiro': get_h('cpf_resp'),
        'email_responsavel': get_h('email_resp_fin') or get_h('email_mae') or get_h('email_pai'),
        'celular_whatsapp_responsavel': phone_resp or phone_mae or phone_pai,
        'cep_endereco': get_h('cep_resp') or get_h('cep_mae') or get_h('cep_pai'),
        'logradouro_endereco': street,
        'numero_endereco': num,
        'complemento_endereco': comp,
        'bairro_endereco': get_h('bairro_resp') or get_h('bairro_mae') or get_h('bairro_pai'),
        'cidade_endereco': get_h('cidade_resp') or get_h('cidade_mae') or 'Indaiatuba',
        'uf_endereco': get_h('uf_resp') or get_h('uf_mae') or 'SP',
        'nome_pai': get_h('nome_pai'),
        'cpf_pai': get_h('cpf_pai'),
        'rg_pai': get_h('rg_pai'),
        'cel_pai': phone_pai,
        'email_pai': get_h('email_pai'),
        'nome_mae': get_h('nome_mae'),
        'cpf_mae': get_h('cpf_mae'),
        'rg_mae': get_h('rg_mae'),
        'cel_mae': phone_mae,
        'email_mae': get_h('email_mae'),
        'is_le_perini': is_le_perini,
        'valor_nominal_anuidade_2027': nominal_2027,
        'tipo_desconto_2027': obs_desc,
        'percentual_desconto_2027': discount_pct,
        'valor_total_anuidade_2027': gross_2027,
        'observacao_desconto_2027': obs_desc,
        'plano_pagamento_anuidade_2027': 13,
        'valor_1a_parcela_2027': first_inst,
        'valor_demais_parcelas_2027': reg_inst,
        'dia_vencimento_parcelas_2027': '10',
        'anuidade_2026_valor_pago': vl_anuid_desc,
        'plano_pagamento_2026': 13,
        'desconto_2026_detalhes': desc_desc,
        'comprador_material_mesmo_responsavel': 'SIM',
        'nome_comprador_material': get_h('Nome_rep_fin'),
        'cpf_comprador_material': get_h('cpf_resp'),
        'valor_total_material': rates['material_total'],
        'valor_total_material_extenso': rates['material_extenso'],
        'numero_parcelas_material': rates['material_parcelas'],
        'valor_parcela_material': rates['material_parc'],
        'valor_parcela_material_extenso': rates['material_parc_extenso'],
        'inicio_vencimento_material': '10 de janeiro de 2027',
        'final_vencimento_material': '10 de dezembro de 2027',
        'dia_vencimento_material': 10,
        'status_rematricula_2027': status_rematricula,
        'status_contrato_escolar_2027': status_contrato,
        'status_contrato_material_2027': status_material,
        'historico_ano_2026': 'Ativo / Em curso',
        'historico_ano_2025': 'Concluído',
        'historico_ano_2024': 'Concluído'
    }
    return record

def main():
    excel_path = r'Planilha de Dados - Colégio Rodin 2026.xlsx'
    print(f"📖 Carregando planilha oficial: {excel_path}...")
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    all_students = []
    rematriculas = []
    concluintes = []
    seen_coc = set()

    legend_keywords = ['falta', 'assinar', 'assinados', 'matriculas novas', 'bolsa', 'resp. finan', 'dra elaine', 'período integral', 'aditamento', 'alunos']

    for sheet_name, (cur_grade, tgt_grade, cycle) in GRADE_PROGRESSION.items():
        if sheet_name not in wb.sheetnames:
            print(f"⚠️ Aba não encontrada: {sheet_name}")
            continue
        ws = wb[sheet_name]
        
        # Mapeamento de cabeçalhos
        header_idx = {}
        for c in range(1, 120):
            val = ws.cell(row=1, column=c).value
            if val:
                header_idx[str(val).strip().lower()] = c

        sheet_count = 0
        for r in range(2, ws.max_row + 1):
            coc_val = ws.cell(row=r, column=header_idx.get('num_coc', 1)).value
            name_val = ws.cell(row=r, column=header_idx.get('nome_aluno', 2)).value
            
            if not coc_val or not name_val:
                continue

            coc_str = str(coc_val).strip()
            name_str = str(name_val).strip()

            if not coc_str or len(name_str) < 3:
                continue

            # Ignorar linhas de legenda/anotação
            if any(k in name_str.lower() for k in legend_keywords):
                continue

            # Deduplicar por RM/COC
            if coc_str in seen_coc:
                print(f"  ⚠️ Duplicata ignorada em {sheet_name}: {name_str} (COC: {coc_str})")
                continue
            seen_coc.add(coc_str)
            
            student_data = parse_student_row(ws, r, header_idx, cur_grade, tgt_grade, cycle)
            if student_data:
                all_students.append(student_data)
                sheet_count += 1
                if tgt_grade == 'Concluinte':
                    concluintes.append(student_data)
                else:
                    rematriculas.append(student_data)

        print(f"  ✅ {sheet_name}: {sheet_count} alunos processados ({cur_grade} -> {tgt_grade})")

    print("\n" + "=" * 60)
    print(f"👥 Total de Alunos Ativos 2026 no Colégio Rodin: {len(all_students)}")
    print(f"📄 Total com Rematrícula 2027 Ativa (7º ano até 3ª série EM): {len(rematriculas)}")
    print(f"🎓 Total Concluintes 2026 (3ª série EM): {len(concluintes)}")
    print("=" * 60)

    # 1. Salvar JSON oficial students2027Data.json
    out_json = 'src/data/students2027Data.json'
    with open(out_json, 'w', encoding='utf-8') as f:
        json.dump(all_students, f, ensure_ascii=False, indent=2)
    print(f"💾 Arquivo gerado: {out_json}")

    # 2. Gerar estatísticas por série de 2027
    grade_counts = {}
    for s in rematriculas:
        g = s['nova_serie_ano_2027']
        grade_counts[g] = grade_counts.get(g, 0) + 1
    
    print("\nDistribuição para Rematrícula 2027:")
    for g, cnt in grade_counts.items():
        print(f"  - {g}: {cnt} alunos")
    print(f"  - Concluintes (sem rematrícula por padrão): {len(concluintes)} alunos")

if __name__ == '__main__':
    main()
