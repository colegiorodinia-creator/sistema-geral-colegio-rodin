import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_raiox_import_template():
    wb = openpyxl.Workbook()

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_code = Font(name="Consolas", size=9, bold=True, color="475569")
    font_data = Font(name="Calibri", size=10, color="0F172A")
    font_data_bold = Font(name="Calibri", size=10, bold=True, color="0F172A")

    # =========================================================================
    # ABA 1: MODELO OFICIAL DE IMPORTAÇÃO RAIO-X (2024 A 2027)
    # =========================================================================
    ws = wb.active
    ws.title = "Importacao_RaioX_2024_2027"
    ws.views.sheetView[0].showGridLines = True

    # Cores dos blocos
    fill_pk = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")       # Vermelho Escuro PK
    fill_cad = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")      # Azul Noturno Cadastro
    fill_2024 = PatternFill(start_color="475569", end_color="475569", fill_type="solid")     # Cinza 2024
    fill_2025 = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")     # Azul Celeste 2025
    fill_2026 = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")     # Verde Petróleo 2026
    fill_2027 = PatternFill(start_color="EA580C", end_color="EA580C", fill_type="solid")     # Laranja Rodin 2027

    bg_pk = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    bg_cad = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    bg_2024 = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    bg_2025 = PatternFill(start_color="F0F9FF", end_color="F0F9FF", fill_type="solid")
    bg_2026 = PatternFill(start_color="F0FDFA", end_color="F0FDFA", fill_type="solid")
    bg_2027 = PatternFill(start_color="FFF7ED", end_color="FFF7ED", fill_type="solid")

    columns = [
        # PK
        {"group": "PRIMARY KEY", "g_fill": bg_pk, "g_color": "991B1B", "fill": fill_pk, "title": "RM (PRIMARY KEY)*", "code": "rm_numero", "w": 20, "align": "center"},

        # CADASTRO DO ESTUDANTE & RESPONSÁVEL
        {"group": "CADASTRO GERAL DO ALUNO E RESPONSÁVEL", "g_fill": bg_cad, "g_color": "1E293B", "fill": fill_cad, "title": "Nome do Aluno*", "code": "nome_aluno", "w": 32, "align": "left"},
        {"group": "CADASTRO GERAL DO ALUNO E RESPONSÁVEL", "g_fill": bg_cad, "g_color": "1E293B", "fill": fill_cad, "title": "CPF Aluno*", "code": "cpf_aluno", "w": 18, "align": "center"},
        {"group": "CADASTRO GERAL DO ALUNO E RESPONSÁVEL", "g_fill": bg_cad, "g_color": "1E293B", "fill": fill_cad, "title": "Data Nasc.*", "code": "data_nascimento", "w": 15, "align": "center"},
        {"group": "CADASTRO GERAL DO ALUNO E RESPONSÁVEL", "g_fill": bg_cad, "g_color": "1E293B", "fill": fill_cad, "title": "Sexo*", "code": "sexo", "w": 10, "align": "center"},
        {"group": "CADASTRO GERAL DO ALUNO E RESPONSÁVEL", "g_fill": bg_cad, "g_color": "1E293B", "fill": fill_cad, "title": "Nome Responsável*", "code": "nome_responsavel", "w": 30, "align": "left"},
        {"group": "CADASTRO GERAL DO ALUNO E RESPONSÁVEL", "g_fill": bg_cad, "g_color": "1E293B", "fill": fill_cad, "title": "CPF Responsável*", "code": "cpf_responsavel", "w": 18, "align": "center"},
        {"group": "CADASTRO GERAL DO ALUNO E RESPONSÁVEL", "g_fill": bg_cad, "g_color": "1E293B", "fill": fill_cad, "title": "WhatsApp Resp.*", "code": "whatsapp_responsavel", "w": 18, "align": "center"},
        {"group": "CADASTRO GERAL DO ALUNO E RESPONSÁVEL", "g_fill": bg_cad, "g_color": "1E293B", "fill": fill_cad, "title": "E-mail Resp.*", "code": "email_responsavel", "w": 26, "align": "left"},
        {"group": "CADASTRO GERAL DO ALUNO E RESPONSÁVEL", "g_fill": bg_cad, "g_color": "1E293B", "fill": fill_cad, "title": "Endereço Completo*", "code": "endereco_completo", "w": 35, "align": "left"},

        # HISTÓRICO 2024
        {"group": "HISTÓRICO 2024", "g_fill": bg_2024, "g_color": "334155", "fill": fill_2024, "title": "Série 2024", "code": "serie_2024", "w": 14, "align": "center"},
        {"group": "HISTÓRICO 2024", "g_fill": bg_2024, "g_color": "334155", "fill": fill_2024, "title": "Turma 2024", "code": "turma_2024", "w": 12, "align": "center"},
        {"group": "HISTÓRICO 2024", "g_fill": bg_2024, "g_color": "334155", "fill": fill_2024, "title": "Turno 2024", "code": "turno_2024", "w": 12, "align": "center"},
        {"group": "HISTÓRICO 2024", "g_fill": bg_2024, "g_color": "334155", "fill": fill_2024, "title": "Anuidade 2024 (R$)", "code": "anuidade_2024", "w": 18, "align": "right"},
        {"group": "HISTÓRICO 2024", "g_fill": bg_2024, "g_color": "334155", "fill": fill_2024, "title": "Contrato Escolar 2024", "code": "status_escolar_2024", "w": 20, "align": "center"},
        {"group": "HISTÓRICO 2024", "g_fill": bg_2024, "g_color": "334155", "fill": fill_2024, "title": "Material 2024", "code": "status_material_2024", "w": 16, "align": "center"},

        # HISTÓRICO 2025
        {"group": "HISTÓRICO 2025", "g_fill": bg_2025, "g_color": "0369A1", "fill": fill_2025, "title": "Série 2025", "code": "serie_2025", "w": 14, "align": "center"},
        {"group": "HISTÓRICO 2025", "g_fill": bg_2025, "g_color": "0369A1", "fill": fill_2025, "title": "Turma 2025", "code": "turma_2025", "w": 12, "align": "center"},
        {"group": "HISTÓRICO 2025", "g_fill": bg_2025, "g_color": "0369A1", "fill": fill_2025, "title": "Turno 2025", "code": "turno_2025", "w": 12, "align": "center"},
        {"group": "HISTÓRICO 2025", "g_fill": bg_2025, "g_color": "0369A1", "fill": fill_2025, "title": "Anuidade 2025 (R$)", "code": "anuidade_2025", "w": 18, "align": "right"},
        {"group": "HISTÓRICO 2025", "g_fill": bg_2025, "g_color": "0369A1", "fill": fill_2025, "title": "Contrato Escolar 2025", "code": "status_escolar_2025", "w": 20, "align": "center"},
        {"group": "HISTÓRICO 2025", "g_fill": bg_2025, "g_color": "0369A1", "fill": fill_2025, "title": "Material 2025", "code": "status_material_2025", "w": 16, "align": "center"},

        # HISTÓRICO 2026
        {"group": "HISTÓRICO 2026", "g_fill": bg_2026, "g_color": "0F766E", "fill": fill_2026, "title": "Série 2026", "code": "serie_2026", "w": 14, "align": "center"},
        {"group": "HISTÓRICO 2026", "g_fill": bg_2026, "g_color": "0F766E", "fill": fill_2026, "title": "Turma 2026", "code": "turma_2026", "w": 12, "align": "center"},
        {"group": "HISTÓRICO 2026", "g_fill": bg_2026, "g_color": "0F766E", "fill": fill_2026, "title": "Turno 2026", "code": "turno_2026", "w": 12, "align": "center"},
        {"group": "HISTÓRICO 2026", "g_fill": bg_2026, "g_color": "0F766E", "fill": fill_2026, "title": "Anuidade 2026 (R$)", "code": "anuidade_2026", "w": 18, "align": "right"},
        {"group": "HISTÓRICO 2026", "g_fill": bg_2026, "g_color": "0F766E", "fill": fill_2026, "title": "Contrato Escolar 2026", "code": "status_escolar_2026", "w": 20, "align": "center"},
        {"group": "HISTÓRICO 2026", "g_fill": bg_2026, "g_color": "0F766E", "fill": fill_2026, "title": "Material 2026", "code": "status_material_2026", "w": 16, "align": "center"},

        # REMATRÍCULA 2027 (ATIVO NO SISTEMA)
        {"group": "REMATRÍCULA 2027 (ANO CORRENTE)", "g_fill": bg_2027, "g_color": "C2410C", "fill": fill_2027, "title": "Nova Série 2027*", "code": "nova_serie_2027", "w": 16, "align": "center"},
        {"group": "REMATRÍCULA 2027 (ANO CORRENTE)", "g_fill": bg_2027, "g_color": "C2410C", "fill": fill_2027, "title": "Turma 2027", "code": "turma_2027", "w": 12, "align": "center"},
        {"group": "REMATRÍCULA 2027 (ANO CORRENTE)", "g_fill": bg_2027, "g_color": "C2410C", "fill": fill_2027, "title": "Turno 2027*", "code": "turno_2027", "w": 12, "align": "center"},
        {"group": "REMATRÍCULA 2027 (ANO CORRENTE)", "g_fill": bg_2027, "g_color": "C2410C", "fill": fill_2027, "title": "Anuidade 2027 (R$)*", "code": "anuidade_2027", "w": 20, "align": "right"},
        {"group": "REMATRÍCULA 2027 (ANO CORRENTE)", "g_fill": bg_2027, "g_color": "C2410C", "fill": fill_2027, "title": "Plano Pgto 2027*", "code": "plano_pagamento_2027", "w": 22, "align": "center"},
        {"group": "REMATRÍCULA 2027 (ANO CORRENTE)", "g_fill": bg_2027, "g_color": "C2410C", "fill": fill_2027, "title": "Valor 1ª Parc 2027 (R$)*", "code": "valor_1a_parcela_2027", "w": 22, "align": "right"},
        {"group": "REMATRÍCULA 2027 (ANO CORRENTE)", "g_fill": bg_2027, "g_color": "C2410C", "fill": fill_2027, "title": "Material 2027 (R$)*", "code": "material_total_2027", "w": 18, "align": "right"},
        {"group": "REMATRÍCULA 2027 (ANO CORRENTE)", "g_fill": bg_2027, "g_color": "C2410C", "fill": fill_2027, "title": "Contrato Anuidade 2027", "code": "status_escolar_2027", "w": 22, "align": "center"},
        {"group": "REMATRÍCULA 2027 (ANO CORRENTE)", "g_fill": bg_2027, "g_color": "C2410C", "fill": fill_2027, "title": "Contrato Material 2027", "code": "status_material_2027", "w": 22, "align": "center"}
    ]

    for idx, col in enumerate(columns, start=1):
        # Linha 2
        c2 = ws.cell(row=2, column=idx, value=col["title"])
        c2.fill = col["fill"]
        c2.font = font_header
        c2.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c2.border = thin_border

        # Linha 3
        c3 = ws.cell(row=3, column=idx, value=col["code"])
        c3.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        c3.font = font_code
        c3.alignment = Alignment(horizontal="center", vertical="center")
        c3.border = thin_border

        col_let = get_column_letter(idx)
        ws.column_dimensions[col_let].width = col["w"]

    # Mesclagem de Grupos na Linha 1
    groups_spec = [
        (1, 1, "PRIMARY KEY", bg_pk, "991B1B"),
        (2, 10, "1. CADASTRO MESTRE DO ESTUDANTE & RESPONSÁVEL", bg_cad, "1E293B"),
        (11, 16, "2. HISTÓRICO 2024 (PASSAGEM PELO COLÉGIO)", bg_2024, "334155"),
        (17, 22, "3. HISTÓRICO 2025 (PASSAGEM PELO COLÉGIO)", bg_2025, "0369A1"),
        (23, 28, "4. HISTÓRICO 2026 (ANO ANTERIOR)", bg_2026, "0F766E"),
        (29, 37, "5. REMATRÍCULA 2027 (ANO CORRENTE ATIVO NO SETOR DE MATRÍCULAS)", bg_2027, "C2410C")
    ]

    for s_c, e_c, title, fill, color in groups_spec:
        if s_c != e_c:
            ws.merge_cells(start_row=1, start_column=s_c, end_row=1, end_column=e_c)
        c = ws.cell(row=1, column=s_c, value=title)
        c.font = Font(name="Calibri", size=11, bold=True, color=color)
        c.alignment = Alignment(horizontal="center", vertical="center")
        for col_idx in range(s_c, e_c + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = fill
            cell.border = thin_border

    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 32
    ws.row_dimensions[3].height = 20

    # 4 Linhas de Exemplo Real
    data_rows = [
        # Christian Pereira (Desde 2024 no Colégio Rodin)
        [
            2553, "Christian Pereira Machado de Campos", "482.910.338-12", "15/04/2012", "Masc.",
            "Natan Machado de Campos Neto", "247.950.638-00", "(19) 99425-0269", "natanneto2000@gmail.com",
            "Homero Paulo Lourenço Barnabé, 345 - Pq São Lourenço, Indaiatuba/SP",
            "6º Ano EF", "A", "Manhã", 25000.00, "Assinado", "Assinado",
            "7º Ano EF", "A", "Manhã", 27500.00, "Assinado", "Assinado",
            "8º Ano EF", "A", "Manhã", 29800.00, "Assinado", "Assinado",
            "9º Ano EF", "A", "Manhã", 34663.20, "13 Parcelas Mensais", 2666.40, 5248.80, "Assinado", "Pendente"
        ],
        # Matheus Milani (Desde 2024 no Colégio Rodin - Vai para o Ensino Médio à Vista)
        [
            2560, "Matheus Milani", "491.829.118-90", "08/11/2011", "Masc.",
            "Wanderson Pedro de Almeida", "57.786.608-4", "(19) 98120-6515", "wpalmeida@hotmail.com",
            "Rua Almerinda Benedita Pacheco de Alcantara, 160 - Jd Jequitibá, Indaiatuba/SP",
            "7º Ano EF", "B", "Manhã", 27500.00, "Assinado", "Assinado",
            "8º Ano EF", "B", "Manhã", 29800.00, "Assinado", "Assinado",
            "9º Ano EF", "B", "Manhã", 31410.19, "Assinado", "Assinado",
            "1ª Série EM", "B", "Manhã", 37752.00, "1x (À Vista com 5% de Desconto)", 35864.40, 5338.20, "Assinado", "Assinado"
        ],
        # Letícia Carli (Entrou em 2025)
        [
            2564, "Letícia Carli Medeiros Silva", "512.441.988-34", "22/09/2014", "Fem.",
            "Yhalle Batista de Lucena", "084.261.824-45", "(19) 98121-9200", "yhalle07@gmail.com",
            "Rua Honduras, 134 - Jardim América, Indaiatuba/SP",
            "—", "—", "—", 0.00, "Não Matriculado", "Não Matriculado",
            "5º Ano EF", "A", "Tarde", 24000.00, "Assinado", "Assinado",
            "6º Ano EF", "A", "Tarde", 26500.00, "Assinado", "Assinado",
            "7º Ano EF", "A", "Tarde", 34663.20, "13 Parcelas Mensais", 2666.40, 5248.80, "Pendente", "Pendente"
        ],
        # Bruno Fialho (Veterano desde 2024)
        [
            2570, "Bruno Fialho de Almeida", "510.866.158-40", "28/03/2015", "Masc.",
            "Wanderson Pedro de Almeida", "57.786.608-4", "(19) 98120-6515", "wpalmeida@hotmail.com",
            "Rua Almerinda Benedita Pacheco de Alcantara, 160 - Jd Jequitibá, Indaiatuba/SP",
            "4º Ano EF", "A", "Manhã", 22000.00, "Assinado", "Assinado",
            "5º Ano EF", "A", "Manhã", 24500.00, "Assinado", "Assinado",
            "6º Ano EF", "A", "Manhã", 27000.00, "Assinado", "Assinado",
            "7º Ano EF", "A", "Manhã", 34663.20, "13 Parcelas Mensais", 2666.40, 5248.80, "Assinado", "Assinado"
        ]
    ]

    for r_i, r_data in enumerate(data_rows, start=4):
        ws.row_dimensions[r_i].height = 24
        for c_i, val in enumerate(r_data, start=1):
            cell = ws.cell(row=r_i, column=c_i, value=val)
            cell.font = font_data
            cell.border = thin_border
            align = columns[c_i-1]["align"]
            cell.alignment = Alignment(horizontal=align, vertical="center")
            if isinstance(val, float) and val > 0:
                cell.number_format = 'R$ #,##0.00'

    # =========================================================================
    # ABA 2: INSTRUÇÕES DE IMPORTAÇÃO E REGRAS DE NEGÓCIO
    # =========================================================================
    ws_inst = wb.create_sheet(title="Instrucoes_Importacao")
    ws_inst.views.sheetView[0].showGridLines = True

    instructions = [
        ("GUIA DE PREENCHIMENTO E IMPORTAÇÃO - RAIO-X DO ALUNO (2024 A 2027)", ""),
        ("", ""),
        ("1. SOBRE A PRIMARY KEY (RM)", ""),
        ("• A coluna 'rm_numero' é a PRIMARY KEY ÚNICA de cada aluno no Colégio Rodin.", ""),
        ("• Cada linha da planilha representa UM ÚNICO ALUNO.", ""),
        ("• O aluno mantém rigorosamente o MESMO RM desde quando ingressou no colégio até o fim dos estudos.", ""),
        ("• Nunca crie RMs diferentes para o mesmo estudante em anos diferentes.", ""),
        ("", ""),
        ("2. HISTÓRICO DE ANOS ANTERIORES (2024, 2025 E 2026)", ""),
        ("• Se o aluno já era matriculado no ano indicado, preencha a série, turma, turno e anuidade praticada.", ""),
        ("• Se o aluno entrou no colégio depois (ex: em 2025 ou 2026), deixe os anos anteriores com '—' e valor 0.00.", ""),
        ("• O sistema reconhecerá automaticamente quantos 'Anos de Casa' o estudante possui no Colégio Rodin.", ""),
        ("", ""),
        ("3. REMATRÍCULA 2027 (ANO ATIVO)", ""),
        ("• Nova Série 2027: a série para a qual o aluno está progredindo (ex: 9º Ano EF, 1ª Série EM).", ""),
        ("• Plano de Pagamento: por padrão '13 Parcelas Mensais' ou '1x (À Vista com 5% de Desconto)'.", ""),
        ("• Status dos Contratos: 'Assinado' ou 'Pendente' para o Requerimento Escolar e Pedido de Material Didático.", ""),
        ("", ""),
        ("4. COMO O SISTEMA LÊ ESSA PLANILHA", ""),
        ("• O sistema processa cada linha, grava a Ficha Cadastral Única no banco e cria a Linha do Tempo da Jornada Escolar.", ""),
        ("• Ao abrir a ficha do aluno no sistema, o setor de matrículas verá o botão 'Raio-X do Aluno', exibindo toda a evolução desde 2024!" "")
    ]

    for idx, item in enumerate(instructions, start=1):
        line = item[0] if isinstance(item, tuple) else item
        cell = ws_inst.cell(row=idx, column=1, value=line)
        if "GUIA" in line or "1." in line or "2." in line or "3." in line or "4." in line:
            cell.font = Font(name="Calibri", size=11, bold=True, color="1E293B")
        else:
            cell.font = Font(name="Calibri", size=10, color="475569")
        ws_inst.row_dimensions[idx].height = 20

    ws_inst.column_dimensions['A'].width = 110

    # Salvar nos dois caminhos
    p1 = r"c:\Users\MARKETING 03\Documents\Antigravity\Sistema Geral Colégio Rodin\public\planilha_importacao_raio_x_colegio_rodin.xlsx"
    p2 = r"c:\Users\MARKETING 03\Documents\Antigravity\Sistema Geral Colégio Rodin\planilha_importacao_raio_x_colegio_rodin.xlsx"

    wb.save(p1)
    wb.save(p2)
    print("Planilha Raio-X gerada com sucesso!")

if __name__ == "__main__":
    create_raiox_import_template()
