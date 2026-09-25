import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_unified_pk_history_template():
    wb = openpyxl.Workbook()

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    font_bold = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_regular = Font(name="Calibri", size=10, color="0F172A")
    font_field_code = Font(name="Consolas", size=9, bold=True, color="64748B")

    # =========================================================================
    # ABA 1: VISÃO UNIFICADA POR RM (1 LINHA POR ALUNO COM HISTÓRICO 24, 25, 26, 27)
    # =========================================================================
    ws1 = wb.active
    ws1.title = "Alunos_Timeline_Por_RM"
    ws1.views.sheetView[0].showGridLines = True

    # Cores dos blocos temporais
    fill_pk = PatternFill(start_color="B91C1C", end_color="B91C1C", fill_type="solid")       # Vermelho Destaque PK
    fill_cad = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")      # Azul Escuro Cadastro
    fill_2024 = PatternFill(start_color="475569", end_color="475569", fill_type="solid")     # Cinza Histórico 24
    fill_2025 = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")     # Azul Claro 25
    fill_2026 = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")     # Verde Água 26
    fill_2027 = PatternFill(start_color="EA580C", end_color="EA580C", fill_type="solid")     # Laranja Rodin 27 (Ativo)

    bg_pk = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    bg_cad = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    bg_2024 = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    bg_2025 = PatternFill(start_color="F0F9FF", end_color="F0F9FF", fill_type="solid")
    bg_2026 = PatternFill(start_color="F0FDFA", end_color="F0FDFA", fill_type="solid")
    bg_2027 = PatternFill(start_color="FFF7ED", end_color="FFF7ED", fill_type="solid")

    cols_ws1 = [
        # PRIMARY KEY
        {"cat": "PRIMARY KEY ÚNICA", "cat_fill": bg_pk, "cat_color": "991B1B", "fill": fill_pk, "label": "RM (PRIMARY KEY)*", "code": "rm_numero", "w": 20, "align": "center"},
        
        # DADOS CADASTRAIS FIXOS
        {"cat": "CADASTRO DO ESTUDANTE & RESPONSÁVEL", "cat_fill": bg_cad, "cat_color": "1E293B", "fill": fill_cad, "label": "Nome do Estudante*", "code": "nome_estudante", "w": 32, "align": "left"},
        {"cat": "CADASTRO DO ESTUDANTE & RESPONSÁVEL", "cat_fill": bg_cad, "cat_color": "1E293B", "fill": fill_cad, "label": "CPF Estudante*", "code": "cpf_estudante", "w": 18, "align": "center"},
        {"cat": "CADASTRO DO ESTUDANTE & RESPONSÁVEL", "cat_fill": bg_cad, "cat_color": "1E293B", "fill": fill_cad, "label": "Data Nascimento*", "code": "data_nasc_estudante", "w": 16, "align": "center"},
        {"cat": "CADASTRO DO ESTUDANTE & RESPONSÁVEL", "cat_fill": bg_cad, "cat_color": "1E293B", "fill": fill_cad, "label": "Nome Responsável*", "code": "nome_responsavel", "w": 30, "align": "left"},
        {"cat": "CADASTRO DO ESTUDANTE & RESPONSÁVEL", "cat_fill": bg_cad, "cat_color": "1E293B", "fill": fill_cad, "label": "CPF Responsável*", "code": "cpf_responsavel", "w": 18, "align": "center"},
        {"cat": "CADASTRO DO ESTUDANTE & RESPONSÁVEL", "cat_fill": bg_cad, "cat_color": "1E293B", "fill": fill_cad, "label": "WhatsApp Resp.*", "code": "whatsapp_responsavel", "w": 18, "align": "center"},
        {"cat": "CADASTRO DO ESTUDANTE & RESPONSÁVEL", "cat_fill": bg_cad, "cat_color": "1E293B", "fill": fill_cad, "label": "E-mail Resp.*", "code": "email_responsavel", "w": 26, "align": "left"},
        {"cat": "CADASTRO DO ESTUDANTE & RESPONSÁVEL", "cat_fill": bg_cad, "cat_color": "1E293B", "fill": fill_cad, "label": "Endereço Completo*", "code": "endereco_completo", "w": 35, "align": "left"},

        # HISTÓRICO 2024
        {"cat": "HISTÓRICO 2024", "cat_fill": bg_2024, "cat_color": "334155", "fill": fill_2024, "label": "Série 2024", "code": "serie_2024", "w": 14, "align": "center"},
        {"cat": "HISTÓRICO 2024", "cat_fill": bg_2024, "cat_color": "334155", "fill": fill_2024, "label": "Turma 2024", "code": "turma_2024", "w": 12, "align": "center"},
        {"cat": "HISTÓRICO 2024", "cat_fill": bg_2024, "cat_color": "334155", "fill": fill_2024, "label": "Anuidade 2024", "code": "anuidade_2024", "w": 18, "align": "right"},
        {"cat": "HISTÓRICO 2024", "cat_fill": bg_2024, "cat_color": "334155", "fill": fill_2024, "label": "Status 2024", "code": "status_2024", "w": 14, "align": "center"},

        # HISTÓRICO 2025
        {"cat": "HISTÓRICO 2025", "cat_fill": bg_2025, "cat_color": "0369A1", "fill": fill_2025, "label": "Série 2025", "code": "serie_2025", "w": 14, "align": "center"},
        {"cat": "HISTÓRICO 2025", "cat_fill": bg_2025, "cat_color": "0369A1", "fill": fill_2025, "label": "Turma 2025", "code": "turma_2025", "w": 12, "align": "center"},
        {"cat": "HISTÓRICO 2025", "cat_fill": bg_2025, "cat_color": "0369A1", "fill": fill_2025, "label": "Anuidade 2025", "code": "anuidade_2025", "w": 18, "align": "right"},
        {"cat": "HISTÓRICO 2025", "cat_fill": bg_2025, "cat_color": "0369A1", "fill": fill_2025, "label": "Status 2025", "code": "status_2025", "w": 14, "align": "center"},

        # HISTÓRICO 2026
        {"cat": "HISTÓRICO 2026", "cat_fill": bg_2026, "cat_color": "0F766E", "fill": fill_2026, "label": "Série 2026", "code": "serie_2026", "w": 14, "align": "center"},
        {"cat": "HISTÓRICO 2026", "cat_fill": bg_2026, "cat_color": "0F766E", "fill": fill_2026, "label": "Turma 2026", "code": "turma_2026", "w": 12, "align": "center"},
        {"cat": "HISTÓRICO 2026", "cat_fill": bg_2026, "cat_color": "0F766E", "fill": fill_2026, "label": "Anuidade 2026", "code": "anuidade_2026", "w": 18, "align": "right"},
        {"cat": "HISTÓRICO 2026", "cat_fill": bg_2026, "cat_color": "0F766E", "fill": fill_2026, "label": "Status 2026", "code": "status_2026", "w": 14, "align": "center"},

        # REMATRÍCULA 2027 (ANO ATUAL)
        {"cat": "REMATRÍCULA 2027 (ATIVO)", "cat_fill": bg_2027, "cat_color": "C2410C", "fill": fill_2027, "label": "Nova Série 2027*", "code": "serie_2027", "w": 16, "align": "center"},
        {"cat": "REMATRÍCULA 2027 (ATIVO)", "cat_fill": bg_2027, "cat_color": "C2410C", "fill": fill_2027, "label": "Turma 2027", "code": "turma_2027", "w": 12, "align": "center"},
        {"cat": "REMATRÍCULA 2027 (ATIVO)", "cat_fill": bg_2027, "cat_color": "C2410C", "fill": fill_2027, "label": "Turno 2027*", "code": "turno_2027", "w": 12, "align": "center"},
        {"cat": "REMATRÍCULA 2027 (ATIVO)", "cat_fill": bg_2027, "cat_color": "C2410C", "fill": fill_2027, "label": "Anuidade 2027 (R$)*", "code": "anuidade_2027", "w": 20, "align": "right"},
        {"cat": "REMATRÍCULA 2027 (ATIVO)", "cat_fill": bg_2027, "cat_color": "C2410C", "fill": fill_2027, "label": "Plano Pgto 2027*", "code": "plano_pgto_2027", "w": 20, "align": "center"},
        {"cat": "REMATRÍCULA 2027 (ATIVO)", "cat_fill": bg_2027, "cat_color": "C2410C", "fill": fill_2027, "label": "Material 2027 (R$)*", "code": "material_2027", "w": 18, "align": "right"},
        {"cat": "REMATRÍCULA 2027 (ATIVO)", "cat_fill": bg_2027, "cat_color": "C2410C", "fill": fill_2027, "label": "Status Contrato Escolar", "code": "status_escolar_2027", "w": 22, "align": "center"},
        {"cat": "REMATRÍCULA 2027 (ATIVO)", "cat_fill": bg_2027, "cat_color": "C2410C", "fill": fill_2027, "label": "Status Material Didático", "code": "status_material_2027", "w": 22, "align": "center"}
    ]

    for idx, c in enumerate(cols_ws1, start=1):
        # Linha 2: Rótulo
        c2 = ws1.cell(row=2, column=idx, value=c["label"])
        c2.fill = c["fill"]
        c2.font = font_bold
        c2.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c2.border = thin_border

        # Linha 3: Código do Campo
        c3 = ws1.cell(row=3, column=idx, value=c["code"])
        c3.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        c3.font = font_field_code
        c3.alignment = Alignment(horizontal="center", vertical="center")
        c3.border = thin_border

        col_let = get_column_letter(idx)
        ws1.column_dimensions[col_let].width = c["w"]

    # Mesclagem das Categorias na Linha 1
    cat_ranges_ws1 = [
        (1, 1, "PRIMARY KEY ÚNICA", bg_pk, "991B1B"),
        (2, 9, "DADOS CADASTRAIS DO ESTUDANTE & RESPONSÁVEIS", bg_cad, "1E293B"),
        (10, 13, "HISTÓRICO 2024 (PASSADO)", bg_2024, "334155"),
        (14, 17, "HISTÓRICO 2025 (PASSADO)", bg_2025, "0369A1"),
        (18, 21, "HISTÓRICO 2026 (ANO ANTERIOR)", bg_2026, "0F766E"),
        (22, 29, "REMATRÍCULA 2027 (ANO ATUAL / EM ANDAMENTO)", bg_2027, "C2410C")
    ]

    for s_c, e_c, title, fill, color in cat_ranges_ws1:
        if s_c != e_c:
            ws1.merge_cells(start_row=1, start_column=s_c, end_row=1, end_column=e_c)
        cell = ws1.cell(row=1, column=s_c, value=title)
        cell.font = Font(name="Calibri", size=11, bold=True, color=color)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        for col_idx in range(s_c, e_c + 1):
            c_cell = ws1.cell(row=1, column=col_idx)
            c_cell.fill = fill
            c_cell.border = thin_border

    ws1.row_dimensions[1].height = 28
    ws1.row_dimensions[2].height = 30
    ws1.row_dimensions[3].height = 20

    # Linhas de Exemplo
    data_examples = [
        # Christian: esteve em 24 (6º), 25 (7º), 26 (8º) e agora 27 (9º)
        [
            2553, "Christian Pereira Machado de Campos", "482.910.338-12", "15/04/2012",
            "Natan Machado de Campos Neto", "247.950.638-00", "(19) 99425-0269", "natanneto2000@gmail.com",
            "Homero Paulo Lourenço Barnabé, 345 - Pq São Lourenço, Indaiatuba/SP",
            "6º Ano EF", "A", 25000.00, "Concluído",
            "7º Ano EF", "A", 27500.00, "Concluído",
            "8º Ano EF", "A", 29800.00, "Concluído",
            "9º Ano EF", "A", "Manhã", 34663.20, "13x de R$ 2.666,40", 5248.80, "Assinado", "Pendente"
        ],
        # Matheus Milani: esteve em 24 (7º), 25 (8º), 26 (9º) e agora 27 (1ª Série EM - À Vista)
        [
            2560, "Matheus Milani", "491.829.118-90", "08/11/2011",
            "Wanderson Pedro de Almeida", "57.786.608-4", "(19) 98120-6515", "wpalmeida@hotmail.com",
            "Rua Almerinda Benedita Pacheco de Alcantara, 160 - Jd Jequitibá, Indaiatuba/SP",
            "7º Ano EF", "B", 27500.00, "Concluído",
            "8º Ano EF", "B", 29800.00, "Concluído",
            "9º Ano EF", "B", 31410.19, "Concluído",
            "1ª Série EM", "B", "Manhã", 37752.00, "1x de R$ 35.864,40 (À Vista)", 5338.20, "Assinado", "Assinado"
        ],
        # Letícia: entrou em 2025 (5º), 2026 (6º) e agora 2027 (7º Ano EF)
        [
            2564, "Letícia Carli Medeiros Silva", "512.441.988-34", "22/09/2014",
            "Yhalle Batista de Lucena", "084.261.824-45", "(19) 98121-9200", "yhalle07@gmail.com",
            "Rua Honduras, 134 - Jardim América, Indaiatuba/SP",
            "—", "—", 0.00, "—",
            "5º Ano EF", "A", 24000.00, "Concluído",
            "6º Ano EF", "A", 26500.00, "Concluído",
            "7º Ano EF", "A", "Tarde", 34663.20, "13x de R$ 2.666,40", 5248.80, "Pendente", "Pendente"
        ]
    ]

    for r_i, r_data in enumerate(data_examples, start=4):
        ws1.row_dimensions[r_i].height = 24
        for c_i, val in enumerate(r_data, start=1):
            cell = ws1.cell(row=r_i, column=c_i, value=val)
            cell.font = font_regular
            cell.border = thin_border
            align = cols_ws1[c_i-1]["align"]
            cell.alignment = Alignment(horizontal=align, vertical="center")
            if isinstance(val, float) and val > 0:
                cell.number_format = 'R$ #,##0.00'

    # =========================================================================
    # ABA 2: ESQUEMA DO BANCO DE DADOS RELACIONAL (SQL PRONTO PARA PRODUÇÃO)
    # =========================================================================
    ws2 = wb.create_sheet(title="Esquema_SQL_Banco_Dados")
    ws2.views.sheetView[0].showGridLines = True

    sql_explanation = [
        ("Como manter a MESMA PRIMARY KEY (RM) sem duplicar alunos e sem estragar a escala do banco de dados:", ""),
        ("", ""),
        ("1. TABELA MESTRE: ESTUDANTES (1 LINHA POR ALUNO - PRIMARY KEY: rm_numero)", ""),
        ("CREATE TABLE estudantes (", ""),
        ("    rm_numero VARCHAR(20) PRIMARY KEY, -- Identificador único da vida inteira do aluno", ""),
        ("    nome_completo VARCHAR(255) NOT NULL,", ""),
        ("    cpf VARCHAR(14) UNIQUE,", ""),
        ("    data_nascimento DATE NOT NULL,", ""),
        ("    sexo VARCHAR(10),", ""),
        ("    naturalidade_cidade VARCHAR(100),", ""),
        ("    naturalidade_uf VARCHAR(2),", ""),
        ("    rg VARCHAR(20),", ""),
        ("    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP", ""),
        (");", ""),
        ("", ""),
        ("2. TABELA DE JORNADA / HISTÓRICO ANUAL (VINCULADA AO MESMO RM - 1 LINHA POR ANO LETIVO)", ""),
        ("CREATE TABLE matriculas_historico (", ""),
        ("    id SERIAL PRIMARY KEY,", ""),
        ("    rm_numero VARCHAR(20) NOT NULL REFERENCES estudantes(rm_numero) ON DELETE RESTRICT,", ""),
        ("    ano_letivo INT NOT NULL, -- 2024, 2025, 2026, 2027...", ""),
        ("    serie_ano VARCHAR(50) NOT NULL, -- 6º Ano EF, 7º Ano EF, etc.", ""),
        ("    turma VARCHAR(10),", ""),
        ("    turno VARCHAR(20),", ""),
        ("    anuidade_total NUMERIC(10,2) NOT NULL,", ""),
        ("    plano_pagamento VARCHAR(50) DEFAULT '13',", ""),
        ("    status_contrato_escolar VARCHAR(30) DEFAULT 'Pendente',", ""),
        ("    status_contrato_material VARCHAR(30) DEFAULT 'Pendente',", ""),
        ("    CONSTRAINT uq_aluno_ano UNIQUE (rm_numero, ano_letivo) -- Garante que o mesmo RM só tem 1 matrícula por ano", ""),
        (");", ""),
        ("", ""),
        ("3. ÍNDICES DE ALTA PERFORMANCE (ESCALA MILIONÁRIA COM RESPOSTA EM 1ms):", ""),
        ("CREATE INDEX idx_matriculas_rm ON matriculas_historico(rm_numero);", ""),
        ("CREATE INDEX idx_matriculas_ano ON matriculas_historico(ano_letivo);", ""),
        ("", ""),
        ("4. COMO CONSULTAR A HISTÓRIA COMPLETA DE UM ALUNO PELO RM (EX: RM 2553):", ""),
        ("SELECT e.rm_numero, e.nome_completo, m.ano_letivo, m.serie_ano, m.anuidade_total, m.status_contrato_escolar", ""),
        ("FROM estudantes e", ""),
        ("JOIN matriculas_historico m ON e.rm_numero = m.rm_numero", ""),
        ("WHERE e.rm_numero = '2553'", ""),
        ("ORDER BY m.ano_letivo ASC; -- Retorna: 2024 -> 2025 -> 2026 -> 2027 em uma única consulta instantânea!", "")
    ]

    for idx, (line, _) in enumerate(sql_explanation, start=1):
        cell = ws2.cell(row=idx, column=1, value=line)
        if "CREATE TABLE" in line or "PRIMARY KEY" in line or "1." in line or "2." in line or "3." in line or "4." in line:
            cell.font = Font(name="Consolas", size=10, bold=True, color="1E293B")
        elif line.startswith("--"):
            cell.font = Font(name="Consolas", size=10, italic=True, color="059669")
        else:
            cell.font = Font(name="Consolas", size=10, color="334155")
        ws2.row_dimensions[idx].height = 18

    ws2.column_dimensions['A'].width = 110

    # Salvar nos dois caminhos oficiais
    path_root = r"c:\Users\MARKETING 03\Documents\Antigravity\Sistema Geral Colégio Rodin\modelo_banco_dados_rematricula_colegio_rodin.xlsx"
    path_public = r"c:\Users\MARKETING 03\Documents\Antigravity\Sistema Geral Colégio Rodin\public\modelo_banco_dados_rematricula_colegio_rodin.xlsx"

    wb.save(path_root)
    wb.save(path_public)
    print("Planilha atualizada com sucesso com visão timeline por RM!")

if __name__ == "__main__":
    generate_unified_pk_history_template()
